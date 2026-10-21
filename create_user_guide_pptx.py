import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ASSETS_DIR = r"c:\Users\COMPUTER\Documents\emr\user_guide_assets"
OUTPUT_PPTX = r"c:\Users\COMPUTER\Documents\emr\Boru_Meda_EMR_User_Guide.pptx"

# Color Palette
NAVY = RGBColor(10, 37, 64)       # #0A2540 Primary Header
TEAL = RGBColor(0, 121, 109)      # #00796D Bahmni Medical Teal
LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC
TEXT_DARK = RGBColor(30, 41, 59)   # #1E293B
TEXT_MUTED = RGBColor(100, 116, 139) # #64748B
WHITE = RGBColor(255, 255, 255)
BORDER_GRAY = RGBColor(226, 232, 240)
ACCENT_BLUE = RGBColor(37, 99, 235) # #2563EB
ACCENT_GREEN = RGBColor(22, 163, 74) # #16A34A
ACCENT_AMBER = RGBColor(217, 119, 6) # #D97706

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def add_header(slide, title, category="BORU MEDA GENERAL HOSPITAL - BAHMNI EMR"):
    # Header bar
    top_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
    top_shape.fill.solid()
    top_shape.fill.fore_color.rgb = NAVY
    top_shape.line.fill.background()

    # Category Pill
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(10), Inches(0.3))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category.upper()
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = RGBColor(45, 212, 191) # Light Teal

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.6))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = WHITE

    # Accent bottom line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.1), Inches(13.333), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()

    # Footer
    footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.15), Inches(11.733), Inches(0.3))
    tf_f = footer.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Boru Meda General Hospital EMR (Server: 196.189.94.149)  |  Standard Operating Procedure & User Manual"
    p_f.font.size = Pt(9)
    p_f.font.color.rgb = TEXT_MUTED

def add_content_slide(slide, title, bullets, image_filename, category="MODULE WALKTHROUGH", note=None):
    add_header(slide, title, category)

    # Left content box
    left_margin = Inches(0.8)
    top_margin = Inches(1.4)
    content_width = Inches(5.8)
    content_height = Inches(5.5)

    # Content background card
    bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_margin, top_margin, content_width, content_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.color.rgb = BORDER_GRAY

    tb = slide.shapes.add_textbox(left_margin + Inches(0.25), top_margin + Inches(0.25), content_width - Inches(0.5), content_height - Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True

    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            header_text, body_text = item
            run_h = p.add_run()
            run_h.text = header_text + ": "
            run_h.font.bold = True
            run_h.font.size = Pt(13)
            run_h.font.color.rgb = TEAL

            run_b = p.add_run()
            run_b.text = body_text
            run_b.font.size = Pt(12)
            run_b.font.color.rgb = TEXT_DARK
        else:
            p.text = item
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(10)

    if note:
        p_note = tf.add_paragraph()
        p_note.text = "TIP / NOTE: " + note
        p_note.font.size = Pt(11)
        p_note.font.italic = True
        p_note.font.bold = True
        p_note.font.color.rgb = ACCENT_AMBER
        p_note.space_before = Pt(8)

    # Right Image box
    img_path = os.path.join(ASSETS_DIR, image_filename)
    if os.path.isfile(img_path):
        img_left = Inches(6.9)
        img_top = Inches(1.4)
        img_width = Inches(5.6)
        img_height = Inches(5.5)

        # Image frame card
        frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, img_left - Inches(0.08), img_top - Inches(0.08), img_width + Inches(0.16), img_height + Inches(0.16))
        frame.fill.solid()
        frame.fill.fore_color.rgb = WHITE
        frame.line.color.rgb = BORDER_GRAY

        # Add image
        slide.shapes.add_picture(img_path, img_left, img_top, width=img_width)

print("[*] Creating Presentation Slides...")

# ── SLIDE 1: TITLE SLIDE ───────────────────────────────────────────────────────
s1 = prs.slides.add_slide(blank_layout)
# Dark Navy background
bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = NAVY
bg1.line.fill.background()

# Teal decorative bar
deco = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.8), Inches(0.15), Inches(3.8))
deco.fill.solid()
deco.fill.fore_color.rgb = TEAL
deco.line.fill.background()

# Title text box
tb1 = s1.shapes.add_textbox(Inches(1.4), Inches(1.8), Inches(10.5), Inches(3.8))
tf1 = tb1.text_frame
tf1.word_wrap = True

p1 = tf1.paragraphs[0]
p1.text = "BORU MEDA GENERAL HOSPITAL"
p1.font.size = Pt(20)
p1.font.bold = True
p1.font.color.rgb = RGBColor(45, 212, 191)

p2 = tf1.add_paragraph()
p2.text = "Comprehensive Bahmni EMR User Guide"
p2.font.size = Pt(36)
p2.font.bold = True
p2.font.color.rgb = WHITE
p2.space_before = Pt(8)

p3 = tf1.add_paragraph()
p3.text = "End-to-End Operational Manual for Clinical, Administrative & Technical Staff"
p3.font.size = Pt(16)
p3.font.color.rgb = RGBColor(203, 213, 225)
p3.space_before = Pt(12)

p4 = tf1.add_paragraph()
p4.text = "Server Endpoint: https://196.189.94.149  |  Version: Bahmni Standard 1.1.0  |  Published: October 2026"
p4.font.size = Pt(12)
p4.font.color.rgb = RGBColor(148, 163, 184)
p4.space_before = Pt(28)

# ── SLIDE 2: SYSTEM ARCHITECTURE & ACCESS ───────────────────────────────────────
s2 = prs.slides.add_slide(blank_layout)
add_content_slide(
    s2,
    "System Access & Architecture Overview",
    [
        ("Server Address", "Access the hospital EMR securely via https://196.189.94.149 over encrypted HTTPS."),
        ("Security Notice", "Self-signed SSL certificate: Click 'Advanced' and proceed to site during first connection."),
        ("Supported Browsers", "Optimized for Google Chrome and Microsoft Edge (desktop viewports 1280x720 or higher)."),
        ("Single Sign-On (SSO)", "One unified credential grants access across Clinical, OpenELIS Laboratory, Odoo Billing, and PACS."),
        ("High Availability", "Powered by containerized Docker services running PostgreSQL, MySQL 8.0, and Apache 2.4.")
    ],
    "01_home_dashboard_bmh.png",
    category="SYSTEM ARCHITECTURE",
    note="Ensure workstations are connected to the internal hospital LAN or authorized health network."
)

# ── SLIDE 3: USER AUTHENTICATION & LOGIN ────────────────────────────────────────
s3 = prs.slides.add_slide(blank_layout)
add_content_slide(
    s3,
    "User Login & Location Session Assignment",
    [
        ("Authentication Credentials", "Enter assigned username (e.g., Admin, Dr. Name, Nurse) and secure password."),
        ("Department Location Picker", "Select your active duty station (e.g., OPD 1, Emergency, General Ward, Lab, Pharmacy)."),
        ("Audit Compliance", "Every login, transaction, and patient record access is permanently stamped with your user ID and location."),
        ("Password Security", "Do not share passwords. Session locks automatically after inactivity to maintain HIPAA/patient privacy."),
        ("Multi-Role Privileges", "System features adapt automatically according to clinical credentials and department roles.")
    ],
    "03_login_page_bmh.png",
    category="AUTHENTICATION & SECURITY",
    note="If your location changes mid-shift, logout and re-select your new ward or clinic room."
)

# ── SLIDE 4: HOME DASHBOARD NAVIGATION ─────────────────────────────────────────
s4 = prs.slides.add_slide(blank_layout)
add_content_slide(
    s4,
    "Home Dashboard & Unified App Launcher",
    [
        ("Central App Grid", "Clean tile-based dashboard providing one-click access to all hospital modules."),
        ("Core Clinical Modules", "Triage, Registration, Programs, Clinical Consultation, InPatient, and OPD Assignment."),
        ("Diagnostic & Ancillary", "Radiology Upload, Patient Documents, Appointment Scheduling, and Reports."),
        ("System Administration", "Admin portal for bulk CSV masterdata, audit trail analysis, bed configs, and backup."),
        ("Consistent Navigation", "Click the 'Home' icon in the top header at any time to return to this dashboard.")
    ],
    "01_home_dashboard_bmh.png",
    category="NAVIGATION & WORKFLOWS",
    note="Apps displayed on your dashboard reflect your role privileges (e.g., Doctors see Clinical; Cashiers see Billing)."
)

# ── SLIDE 5: PATIENT REGISTRATION & DEMOGRAPHICS ────────────────────────────────
s5 = prs.slides.add_slide(blank_layout)
add_content_slide(
    s5,
    "Patient Registration & Identity Management",
    [
        ("Unique Patient Identifier (MRN)", "System automatically generates unique Medical Record Numbers for every patient."),
        ("Duplicate Prevention", "Search existing patients by Name, Phone Number, or National ID before creating new records."),
        ("Ethiopian Calendar Support", "Built-in dual Ethiopian & Gregorian calendar converter for birthdates and visit dates."),
        ("Demographic Capturing", "Record Woreda, Kebele, House No, Contact Info, Emergency Contact, and Sponsoring Entity."),
        ("Direct Action Routing", "Quickly start an OPD Visit, Emergency Triage Visit, or direct Consultation from registration.")
    ],
    "08_registration_form_bmh.png",
    category="PATIENT ADMINISTRATION",
    note="Always verify Phone Number and Kebele to facilitate emergency contact and follow-up tracing."
)

# ── SLIDE 6: TRIAGE & EMERGENCY ASSESSMENT ──────────────────────────────────────
s6 = prs.slides.add_slide(blank_layout)
add_content_slide(
    s6,
    "Triage & Vital Signs Recording",
    [
        ("Vital Signs Entry", "Record Blood Pressure (Systolic/Diastolic), Pulse, Respiratory Rate, Temp, and SpO2."),
        ("Emergency Oxygen Tracking", "Indicate oxygen therapy requirements and flow rates for critical patient monitoring."),
        ("Acuity Classification", "Tag patient condition: Red (Emergency), Yellow (Urgent), or Green (Stable non-urgent)."),
        ("Automated Flowsheet", "Vitals entered at triage instantly appear on the doctor's consultation dashboard in real time."),
        ("Fast-Track Routing", "Immediately transfers patient queue to OPD doctor or emergency physician without delay.")
    ],
    "06_triage_clinical_bmh.png",
    category="EMERGENCY & TRIAGE",
    note="Critical triage alerts flash on clinical workstations to ensure immediate emergency response."
)

# ── SLIDE 7: OUTPATIENT CLINICAL CONSULTATION ──────────────────────────────────
s7 = prs.slides.add_slide(blank_layout)
add_content_slide(
    s7,
    "Outpatient Clinical Consultation & EHR",
    [
        ("Holistic Patient Summary", "View complete medical history, previous visits, allergies, chronic conditions, and past vitals."),
        ("Chief Complaint & HPI", "Document narrative clinical findings, history of present illness, and physical exam findings."),
        ("ICD-10 / CIEL Diagnosis", "Search standardized diagnoses with coded ICD-10 mapping (Primary & Secondary diagnoses)."),
        ("Clinical Dispositions", "Admit to Ward, Discharge, Transfer, Follow-up Appointment, or Inter-Facility Referral."),
        ("Digital Signature", "Consultation note automatically captures attending doctor name and timestamp.")
    ],
    "09_consultation_dashboard_bmh.png",
    category="CLINICAL EHR",
    note="Use the 'Save' button regularly during long examinations to persist clinical notes safely."
)

# ── SLIDE 8: ORDERS, LAB & INVESTIGATIONS ───────────────────────────────────────
s8 = prs.slides.add_slide(blank_layout)
add_content_slide(
    s8,
    "Investigations, Orders & Lab Sync",
    [
        ("Order Selector", "Search lab tests (CBC, Urinalysis, Blood Chemistry) and radiology exams (X-Ray, Ultrasound, CT)."),
        ("Order Templates", "Pre-configured bundles for Antenatal, Trauma, Diabetic, and Fever workups save ordering time."),
        ("OpenELIS Bidirectional Sync", "Orders transfer electronically to OpenELIS Laboratory; results stream back automatically."),
        ("Urgency Levels", "Flag urgent tests as STAT / Emergency for prioritized laboratory processing."),
        ("Cumulative Flowsheets", "Track lab parameter trends over time across consecutive patient visits.")
    ],
    "10_orders_investigations_bmh.png",
    category="INVESTIGATIONS & LAB",
    note="Doctors can review lab results directly on the patient dashboard without waiting for paper slips."
)

# ── SLIDE 9: INPATIENT ADMISSION & BED MANAGEMENT ──────────────────────────────
s9 = prs.slides.add_slide(blank_layout)
add_content_slide(
    s9,
    "Inpatient Admission & Bed Management",
    [
        ("Visual Ward Map", "Interactive graphical layout of all hospital wards (Medical, Surgical, Pediatric, Maternity, ICU)."),
        ("Real-Time Bed Status", "Color-coded beds show Available (Green), Occupied (Blue), and Maintenance (Gray)."),
        ("Seamless Bed Allocation", "Drag-and-drop or single-click admission directly assigns admitted patient to vacant bed."),
        ("Inter-Ward Transfer", "Move patients between beds or wards with automatic audit history and nurse handoff."),
        ("Discharge Processing", "Frees the bed instantly upon clinical discharge and finalizes inpatient stay billing.")
    ],
    "11_bed_management_bmh.png",
    category="INPATIENT ADMISSIONS",
    note="Bed occupancy rates update dynamically, helping nursing supervisors manage ward capacity."
)

# ── SLIDE 10: REFERRAL MANAGEMENT & PRINTABLE SLIP ──────────────────────────────
s10 = prs.slides.add_slide(blank_layout)
add_content_slide(
    s10,
    "Inter-Facility Referral & Printable Referral Slip",
    [
        ("Referral Disposition", "Select 'Refer Patient' disposition in consultation to initiate transfer workflow."),
        ("Comprehensive Transfer Record", "Captures Receiving Hospital, Specialty Dept, Reason, Clinical History, and Treatments."),
        ("Synced Facility Branding", "Referral Slip dynamically prints 'BORU MEDA GENERAL HOSPITAL' and official logo."),
        ("Ambulance & Escort Logistics", "Flags requirements for medical escorts, oxygen support, and ambulance dispatch."),
        ("Standardized Print Layout", "Generates clean, standardized A4 referral slip ready for receiving facility handoff.")
    ],
    "13_referral_form_print_bmh.png",
    category="PATIENT REFERRAL WORKFLOW",
    note="Click 'Print Referral Form' to generate an official hard copy for ambulance transport."
)

# ── SLIDE 11: PATIENT DOCUMENTS & RADIOLOGY UPLOADS ─────────────────────────────
s11 = prs.slides.add_slide(blank_layout)
add_content_slide(
    s11,
    "Digital Document & Radiology Scan Uploads",
    [
        ("Paperless Record Keeping", "Upload external referral letters, scanned lab sheets, identity cards, and discharge notes."),
        ("Radiology Upload Module", "Attach X-Ray, CT Scan, MRI, and Ultrasound digital images directly to patient file."),
        ("Multi-Format Compatibility", "Accepts JPEG, PNG, PDF, and DICOM imaging formats up to 50MB per upload."),
        ("Encounter Integration", "Uploaded files become permanent clinical attachments visible on the patient dashboard."),
        ("Secure Docker Volume Storage", "All files reside on dedicated encrypted storage volumes backed up automatically.")
    ],
    "04_patient_documents_bmh.png",
    category="DOCUMENT ARCHIVING",
    note="Clearly label uploaded document titles (e.g., 'Chest X-Ray 21-Oct-2026') for rapid clinical retrieval."
)

# ── SLIDE 12: SYSTEM ADMINISTRATION PORTAL ──────────────────────────────────────
s12 = prs.slides.add_slide(blank_layout)
add_content_slide(
    s12,
    "Hospital Administration & Master Data Control",
    [
        ("CSV Bulk Upload", "Batch import patient registries, ICD-10 concepts, and pharmaceutical inventory via CSV."),
        ("CSV Export Tools", "Export demographic cohorts, disease reports, and operational metrics for health bureau analysis."),
        ("Audit Log Examiner", "Inspect detailed system event logs, user activities, timestamped edits, and logins."),
        ("Order Set Dashboard", "Create and maintain clinical order sets, custom drug protocols, and test bundles."),
        ("Bed Configuration", "Add, rename, or reconfigure hospital wards, rooms, and bed numbers as facility expands.")
    ],
    "02_admin_dashboard_bmh.png",
    category="ADMINISTRATION",
    note="Administrative functions are restricted to authorized IT officers and Medical Directors."
)

# ── SLIDE 13: SYSTEM BACKUP & DISASTER RECOVERY ─────────────────────────────────
s13 = prs.slides.add_slide(blank_layout)
add_content_slide(
    s13,
    "Disaster Recovery: System Backup & Restore",
    [
        ("Integrated Admin Tile", "Accessible directly inside the Admin portal under 'Backup & Restore'."),
        ("Full System Backup", "Archives all databases (OpenMRS, OpenELIS, Odoo, Reports, PACS), 11 Docker volumes, and configs."),
        ("Database-Only Backup", "Lightweight, high-speed SQL dump of all relational databases for daily automated archiving."),
        ("Live Progress & Logs Console", "Visual progress bar and live terminal stream operations in real time without page refresh."),
        ("One-Click Disaster Restore", "Restores system from selected archive with automated UI patch protection and safe restart.")
    ],
    "14_backup_restore_app.png",
    category="DISASTER RECOVERY",
    note="Backups are stored locally in /opt/emr/backups/ and automatically pruned with a 10-archive retention limit."
)

# ── SLIDE 14: MAINTENANCE & IT SUPPORT ──────────────────────────────────────────
s14 = prs.slides.add_slide(blank_layout)
add_content_slide(
    s14,
    "Routine Maintenance & Helpdesk Procedures",
    [
        ("Daily Automated Backups", "Scheduled nightly backups safeguard patient data without interrupting shift operations."),
        ("Zero-Downtime Patching", "UI updates and configurations deploy dynamically without restarting clinical containers."),
        ("Troubleshooting Connectivity", "Verify LAN network cable, ensure IP is 196.189.94.149, and clear browser cache if needed."),
        ("Password Resets", "Hospital IT administrator can reset credentials instantly via the OpenMRS administration console."),
        ("Technical Escalation", "Contact Boru Meda Hospital IT Support Desk for urgent assistance or hardware support.")
    ],
    "01_home_dashboard_bmh.png",
    category="SUPPORT & SOP",
    note="Hospital IT Helpdesk: Contact Hospital Liaison Office / IT Department on Ext. 101."
)

prs.save(OUTPUT_PPTX)
print(f"[✔] PowerPoint saved: {OUTPUT_PPTX} ({os.path.getsize(OUTPUT_PPTX)} bytes)")
